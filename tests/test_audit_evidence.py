import copy
import json
from pathlib import Path
import subprocess
import sys

import pytest

from tools.provenance import artifact, atomic_json, build_manifest, validate_manifest, sha256_file
from tools.audit_training_provenance import audit
from tools.create_quality_control_workbook import create
from tools.audit_quality_control_workbook import audit as audit_qc
from test_training_plan import prepared_fixture

ROOT=Path(__file__).resolve().parents[1]

def inputs(tmp_path):
    data=tmp_path/'data'; data.mkdir()
    for name in ('train_00000.parquet','train_00001.parquet','val_00000.parquet'):
        (data/name).write_bytes(name.encode())
    tok=tmp_path/'tokenizer.pkl';tok.write_bytes(b'tokenizer fixture')
    config=tmp_path/'config.json';config.write_text('{"num_iterations": 7}')
    return data,tok,config

def test_manifest_cli_binds_contents_and_does_not_overwrite(tmp_path):
    data,tok,config=inputs(tmp_path)
    out=tmp_path/'run.json'
    args=[sys.executable,str(ROOT/'tools/write_training_run_manifest.py'),'--model-tag','test',
          '--phase','base','--dataset-dir',str(data),'--tokenizer',str(tok),
          '--config-json',str(config),'--output',str(out)]
    assert subprocess.run(args,capture_output=True).returncode==0
    record=json.loads(out.read_text());validate_manifest(record)
    assert record['status']=='PREPARED' and record['checkpoints']==[]
    assert len(record['dataset']['train'])==2
    assert subprocess.run(args,capture_output=True).returncode!=0
    (data/'train_00000.parquet').write_bytes(b'changed')
    with pytest.raises(ValueError,match='changed'): validate_manifest(record)

def test_provenance_rejects_absent_unrelated_and_changed_evidence(tmp_path):
    assert audit(tmp_path)['audit_status']=='FAIL'
    mdir=tmp_path/'run_manifests';mdir.mkdir()
    (mdir/'bogus_RUN_MANIFEST.json').write_text('{}')
    assert audit(tmp_path)['audit_status']=='FAIL'
    (mdir/'bogus_RUN_MANIFEST.json').unlink()
    data,tok,_=inputs(tmp_path)
    ck=tmp_path/'base_checkpoints/test/model_000001.pt';ck.parent.mkdir(parents=True);ck.write_bytes(b'weights')
    record=build_manifest(model_tag='test',phase='base',config={'iterations':1},dataset_dir=data,tokenizer_path=tok,checkpoint_path=ck)
    atomic_json(mdir/'test_RUN_MANIFEST.json',record)
    assert audit(tmp_path)['audit_status']=='PASS'
    (data/'train_00002.parquet').write_bytes(b'new shard')
    assert audit(tmp_path)['audit_status']=='FAIL'

def qc_inputs(pack):
    (pack/'eval').mkdir()
    for name,strict in [('strict_holdout_quality_control_v2.be.jsonl',True),('regression_quality_control_v1.be.jsonl',False)]:
        (pack/'eval'/name).write_text(json.dumps({'eval_id':name,'prompt':'Правер адказ.','never_train':strict})+'\n')

def test_qc_absence_is_unknown_and_failed_checks_cannot_pass(tmp_path):
    pytest.importorskip('openpyxl')
    from openpyxl import load_workbook
    qc_inputs(tmp_path)
    report=create(tmp_path)
    assert set(report['checks'].values())=={'UNKNOWN'}
    path=tmp_path/'reports/eval/BELKA_QUALITY_CONTROL.xlsx'
    assert audit_qc(tmp_path,path)['audit_status']=='FAIL'
    wb=load_workbook(path)
    wb['Strict_Holdout_LOCKED']['D2']=False
    wb['Training_Gates']['B3']='YES'
    wb.save(path)
    failure=audit_qc(tmp_path,path)
    assert failure['audit_status']=='FAIL'
    assert any('never_train mismatch' in x for x in failure['failures'])
    assert 'training/evidence gates changed' in failure['failures']

def test_qc_detects_workbook_prompt_tampering(tmp_path):
    pytest.importorskip('openpyxl')
    from openpyxl import load_workbook
    qc_inputs(tmp_path);create(tmp_path)
    path=tmp_path/'reports/eval/BELKA_QUALITY_CONTROL.xlsx'
    wb=load_workbook(path);wb['Strict_Holdout_LOCKED']['B2']='changed';wb.save(path)
    assert any('exact prompt identity' in x for x in audit_qc(tmp_path,path)['failures'])

def test_canonical_sft_manifest_and_preparation_are_not_checkpoint_proof(tmp_path):
    data,tok,_=inputs(tmp_path)
    (data/'identity_conversations.jsonl').write_text('{"messages":[]}\n')
    (data/'identity_conversations_val.jsonl').write_text('{"messages":[]}\n')
    record=build_manifest(model_tag='test',phase='sft',config={'iterations':1},dataset_dir=data,tokenizer_path=tok)
    validate_manifest(record)
    assert Path(record['dataset']['train'][0]['path']).name=='identity_conversations.jsonl'
    atomic_json(tmp_path/'run_manifests/prepared_RUN_MANIFEST.json',record)
    report=audit(tmp_path)
    assert report['audit_status']=='FAIL' and len(report['prepared_manifests'])==1

def test_qc_valid_subset_evidence_then_changed_dashboard_and_sft_fail(prepared_fixture, monkeypatch):
    from openpyxl import load_workbook
    import data_pipeline.h200_admission as admission
    import data_pipeline.h200_evidence as corpus_evidence
    from test_h200_subset import select_parent
    from tools.subset_h200_data import subset
    base, _ = prepared_fixture
    parent = select_parent(base)
    parent_report = parent/'H200_CORPUS_MANIFEST.json'
    original_parent = json.loads(parent_report.read_text())
    original_parent['holdout']['scope'] = 'fixture '+('x'*40000)
    parent_report.write_text(json.dumps(original_parent))
    # Produce a real exact-row subset whose counters differ from the parent.
    monkeypatch.setattr(admission, 'admission_reason',
        lambda row, policy: 'fixture_rejection' if row['input_row'] == 0 else None)
    proof_path = base/'QC_DATA_PREPARATION.json'
    proof = subset(base, proof_path, shard_rows=3)
    data = Path(proof['corpus_dir']); sft = (base/'.sft_current').resolve()
    assert proof['splits']['train']['rows'] == 7
    assert proof['holdout']['checked_train_rows'] == 8
    assert proof['benchmark_decontamination']['checked_final_rows'] == 16
    ck = base/'base_checkpoints/test/model_000001.pt'; ck.parent.mkdir(parents=True); ck.write_bytes(b'weights')
    manifest = build_manifest(model_tag='test', phase='base', config={'iterations': 1},
        dataset_dir=data, tokenizer_path=base/'tokenizer/tokenizer.pkl', checkpoint_path=ck)
    atomic_json(base/'run_manifests/run_RUN_MANIFEST.json', manifest)
    provenance_path = base/'training_provenance_audit.json'
    atomic_json(provenance_path, audit(base))
    path = base/'BELKA_QUALITY_CONTROL.xlsx'
    def regenerate(): return create(ROOT, path, proof_path, provenance_path)
    result = regenerate()
    assert result['checks'] == {'corpus': 'PASS', 'leakage': 'PASS', 'provenance': 'PASS'}
    assert result['subset_derivation'] == proof['subset_derivation']
    assert result['measurements']['evidence_mode'] == 'inherited_whole_row_subset'
    assert result['measurements']['current_train_rows'] == 7
    assert result['measurements']['parent_checked_train_rows'] == 8
    assert audit_qc(ROOT, path)['audit_status'] == 'PASS'
    wb = load_workbook(path); wb['Regression_Dashboard']['B2'] = 999; wb.save(path)
    assert 'dashboard measurements changed' in audit_qc(ROOT, path)['failures']
    regenerate()
    original_hash = corpus_evidence.sha256_file
    with monkeypatch.context() as changed_policy:
        changed_policy.setattr(corpus_evidence, 'sha256_file', lambda p:
            '0'*64 if Path(p) == ROOT/'configs/h200_admission_policy.json' else original_hash(p))
        assert audit_qc(ROOT, path)['audit_status'] == 'FAIL'
    assert audit_qc(ROOT, path)['audit_status'] == 'PASS'
    for name in ('quality_decisions.jsonl', 'subset_decisions.jsonl', 'PARENT_H200_CORPUS_MANIFEST.json'):
        ledger = data/name; original = ledger.read_bytes(); ledger.write_bytes(original+b' ')
        assert audit_qc(ROOT, path)['audit_status'] == 'FAIL'
        ledger.write_bytes(original)
        assert audit_qc(ROOT, path)['audit_status'] == 'PASS'
    bm = proof['benchmark_decontamination']
    for field, value in [('manifest_sha256', '0'*64), ('checked_final_rows', 14), ('checked_sft_messages', 0),
                         ('remaining_corpus_hits', 1), ('remaining_sft_hits', 1), ('files', bm['files'][:-1]),
                         ('coverage', dict(records=1))]:
        changed = copy.deepcopy(proof); changed['benchmark_decontamination'][field] = value
        proof_path.write_text(json.dumps(changed))
        fresh = regenerate()
        assert fresh['checks']['leakage'] == 'FAIL' and any('benchmark' in issue for issue in fresh['problems'])
        assert audit_qc(ROOT, path)['audit_status'] == 'FAIL'
    proof_path.write_text(json.dumps(proof)); regenerate()
    assert audit_qc(ROOT, path)['audit_status'] == 'PASS'
    (sft/'identity_conversations.jsonl').write_text('changed')
    assert audit_qc(ROOT, path)['audit_status'] == 'FAIL'


def test_qc_accepts_portable_subset_lineage_but_no_checkpoint_claim(prepared_fixture):
    from tools.transfer_training_data import export_data, restore_data
    base, _ = prepared_fixture; archive = base/'qc-transfer.tar'
    receipt = export_data(base, archive); restored = base/'qc-server-copy'
    restore_data(archive, restored, receipt['sha256'])
    corpus = (restored/'.corpus_current').resolve()
    result = create(ROOT, base/'restored.xlsx', corpus/'H200_CORPUS_MANIFEST.json', base/'absent-provenance.json')
    assert result['checks'] == {'corpus': 'PASS', 'leakage': 'PASS', 'provenance': 'UNKNOWN'}
    assert all(Path(f['path']).is_relative_to(restored) for f in result['dataset_sft'])
    assert result['measurements']['evidence_mode'] == 'inherited_whole_row_subset'
    assert audit_qc(ROOT, base/'restored.xlsx')['audit_status'] == 'FAIL'

def test_integrity_absent_empty_and_populated_books(tmp_path,monkeypatch):
    import tools.audit_repo_integrity as integrity
    monkeypatch.setattr(integrity,'PACK_DIR',tmp_path)
    assert integrity.dir_file_count('books')==-1
    (tmp_path/'books').mkdir()
    assert integrity.dir_file_count('books')==0
    (tmp_path/'books/a.txt').write_text('book')
    assert integrity.dir_file_count('books')==1
