#!/usr/bin/env python3
"""Fail closed on missing, changed or unverified checkpoint QC evidence."""
from __future__ import annotations
import argparse
import json
import os
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from tools.create_quality_control_workbook import evidence, prompts, MODEL_RUNS_COLS, display_measurement
from tools.provenance import atomic_json, verify_artifact

def audit(pack, path):
    from openpyxl import load_workbook
    failures=[]
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        required={'README','Strict_Holdout_LOCKED','Regression_SeenIntent','Model_Runs',
                  'Strict_Holdout_Results','Regression_Results','Regression_Dashboard',
                  'Dataset_Registry','Training_Gates','Source_Evidence'}
        if set(wb.sheetnames)!=required:
            raise ValueError('missing or unexpected QC sheets')
        saved=json.loads(''.join(row[0] for row in wb['Source_Evidence'].iter_rows(min_row=2,values_only=True)))
        for record in saved['inputs'].values():
            verify_artifact(record)
        for record in saved['reports'].values():
            verify_artifact(record)
        current=evidence(pack, saved['reports'].get('data',{}).get('path'),
                         saved['reports'].get('provenance',{}).get('path'))
        if current!=saved:
            failures.append('evidence has changed since workbook generation')
        def rows(name):
            return list(wb[name].iter_rows(min_row=2,values_only=True))
        for name,key,strict in [('Strict_Holdout_LOCKED','strict',True),('Regression_SeenIntent','regression',False)]:
            expected=prompts(current['inputs'][key]['path'],strict=strict)
            if rows(name)!=expected:
                failures.append(f'{name}: exact prompt identity/content/never_train mismatch')
        expected_gates=[('TRAINING_ALLOWED','NO'),('SFT_ALLOWED','NO')]+list(current['checks'].items())
        if rows('Training_Gates')!=expected_gates:
            failures.append('training/evidence gates changed')
        for name,status in current['checks'].items():
            if status!='PASS':
                failures.append(f'{name}: {status}')
        if list(next(wb['Model_Runs'].iter_rows(values_only=True)))!=MODEL_RUNS_COLS:
            failures.append('model provenance columns changed')
        if rows('Model_Runs')!=[tuple(row) for row in current['model_runs']] or not current['model_runs']:
            failures.append('model provenance missing or changed')
        expected_data=[tuple(f[k] for k in ('path','sha256','rows','split')) for f in current['dataset']]
        if not expected_data or rows('Dataset_Registry')!=expected_data:
            failures.append('dataset registry missing or changed')
        expected_dashboard=[('strict_prompts',len(prompts(current['inputs']['strict']['path'],strict=True))),
                            ('regression_prompts',len(prompts(current['inputs']['regression']['path'])))]+[
                            (k,display_measurement(v)) for k,v in current['measurements'].items()]
        if rows('Regression_Dashboard')!=expected_dashboard:
            failures.append('dashboard measurements changed')
        for name in ('Strict_Holdout_Results','Regression_Results'):
            if rows(name)!=[('NOT_EVALUATED',)]:
                failures.append(f'{name}: unsupported quality claim')
    except (OSError,ValueError,TypeError,KeyError) as exc:
        failures.append(str(exc))
    finally:
        wb.close()
    return dict(schema_version=1,audit_status='FAIL' if failures else 'PASS',
                scope='checkpoint QC evidence, not model quality or training permission',
                TRAINING_ALLOWED='NO',SFT_ALLOWED='NO',failures=failures)

def main(argv=None):
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--pack-dir',type=Path,default=Path(os.environ.get('PACK_DIR','.')))
    ap.add_argument('--workbook',type=Path)
    ap.add_argument('--output',type=Path)
    args=ap.parse_args(argv)
    try:
        report=audit(args.pack_dir,args.workbook or args.pack_dir/'reports/eval/BELKA_QUALITY_CONTROL.xlsx')
    except (OSError,ValueError,ImportError) as exc:
        ap.error(str(exc))
    atomic_json(args.output or args.pack_dir/'reports/eval/QUALITY_CONTROL_XLSX_AUDIT.json',report,overwrite=True)
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return int(report['audit_status']!='PASS')

if __name__=='__main__':
    raise SystemExit(main())
